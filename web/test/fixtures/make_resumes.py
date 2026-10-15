"""Builds the resume fixtures for the Lite resume-reader tests (web/test/resume.test.ts).

Every person here is invented. The layouts copy formats people really use: a modern template with spaced-out
headings and font-size headers, the LaTeX "Jake's Resume" template, a classic one-column resume, a two-column
sidebar resume, a Word document and plain text.

Run from the repo root:  api/.venv/Scripts/python web/test/fixtures/make_resumes.py
"""
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

OUT = Path(__file__).parent / "resumes"
OUT.mkdir(exist_ok=True)
W, H = A4
R, B, I = "Helvetica", "Helvetica-Bold", "Helvetica-Oblique"


def text(c, x, y, s, font=R, size=9.2):
    c.setFont(font, size)
    c.drawString(x, y, s)


def right(c, x, y, s, font=R, size=7.8):
    c.setFont(font, size)
    c.drawRightString(x, y, s)


# ------------------------------------------------------------------ 1. modern: spaced headings, size-based headers
def modern():
    c = canvas.Canvas(str(OUT / "modern-spaced.pdf"), pagesize=A4)
    text(c, 42.5, 767.7, "Ananya Iyer", B, 27)
    text(c, 42.5, 751.2, "D A T A", R, 8)
    text(c, 42.5 + c.stringWidth("D A T A", R, 8) + 7, 751.2, "S C I E N C E", R, 8)
    for y, s in [(801.4, "Bengaluru, India"), (789.4, "ananya.iyer.dev@example.com"), (776.7, "linkedin.com/in/ananya-iyer-dev"), (764.7, "github.com/ananya-iyer")]:
        right(c, 552.7, y, s)
    text(c, 42.5, 722.7, "Final-year engineering student who builds data products end to end, from SQL pipelines to models that", R, 9.4)
    text(c, 42.5, 709.2, "ship. Two internships in fintech analytics and a habit of measuring everything.", R, 9.4)

    y = 680
    text(c, 42.5, y, "E X P E R I E N C E", R, 7.8)
    jobs = [
        ("Razorpay", "Data Science Intern · Remote", "May 2025 – Aug 2025", "Payments fraud analytics",
         [["Built a gradient boosting model that flagged 2,300 risky merchants a month, improving precision by",
           "18% over the rules engine."],
          ["Automated weekly risk reporting in SQL and Python, saving the team 5 hours a week."]],
         "Python · SQL · XGBoost · Airflow"),
        ("Zeta Suite", "Analytics Intern", "Dec 2024 – Feb 2025", None,
         [["Cleaned and joined 12 million card transactions to find the top reasons for failed payments."],
          ["Presented a Tableau dashboard to the product team; two fixes shipped from it."]],
         "SQL · Tableau"),
    ]
    y -= 18
    for company, title, dates, sub, paras, tech in jobs:
        text(c, 42.5, y, company, B, 10.6)
        text(c, 42.5 + c.stringWidth(company, B, 10.6) + 5, y, title, R, 9)
        right(c, 552.7, y, dates)
        y -= 14
        if sub:
            text(c, 42.5, y, sub, R, 8.8)
            y -= 14
        for para in paras:
            for line in para:
                text(c, 53.8, y, line)
                y -= 13
        text(c, 42.5, y, tech, R, 7.6)
        y -= 22

    text(c, 42.5, y, "P R O J E C T S", R, 7.8)
    y -= 18
    projects = [
        ("Churn Radar", "Customer churn early-warning dashboard · personal", "2025",
         ["Predicts which subscribers will cancel in the next 30 days from usage logs, with an", "AUC of 0.87 on held-out data."]),
        ("ResumeLens", "NLP resume screener · hackathon", "Mar 2024",
         ["Ranks resumes against a job description with sentence embeddings; won 2nd place among 60 teams."]),
    ]
    for name, sub, dates, lines in projects:
        text(c, 42.5, y, name, B, 10.6)
        text(c, 42.5 + c.stringWidth(name, B, 10.6) + 5, y, sub, R, 9)
        right(c, 552.7, y, dates)
        y -= 14
        for line in lines:
            text(c, 53.8, y, line)
            y -= 13
        y -= 9

    text(c, 42.5, y, "E D U C A T I O N", R, 7.8)
    y -= 18
    text(c, 42.5, y, "RV College of Engineering", B, 10.6)
    text(c, 42.5 + c.stringWidth("RV College of Engineering", B, 10.6) + 5, y, "B.E., Computer Science · Bengaluru", R, 9)
    right(c, 552.7, y, "2022 – 2026")
    y -= 14
    text(c, 42.5, y, "CGPA: 8.9/10", R, 8.8)
    y -= 24

    text(c, 42.5, y, "S K I L L S", R, 7.8)
    y -= 16
    for cat, items in [("Languages", "Python, SQL, R"), ("ML", "scikit-learn, XGBoost, PyTorch, NLP"), ("Tools", "Airflow, Tableau, Git, Docker")]:
        text(c, 42.5, y, cat, B, 8.8)
        text(c, 120, y, items, R, 8.8)
        y -= 13
    c.save()


# ------------------------------------------------------------------ 2. Jake's Resume (LaTeX)
def jakes():
    c = canvas.Canvas(str(OUT / "jakes-template.pdf"), pagesize=A4)
    c.setFont(B, 24)
    c.drawCentredString(W / 2, 790, "Rohan Mehta")
    c.setFont(R, 9)
    c.drawCentredString(W / 2, 772, "+91 98200 12345 | rohan.mehta@example.com | linkedin.com/in/rohanmehta | github.com/rohanmehta")
    y = 748

    def heading(label):
        nonlocal y
        text(c, 40, y, label, R, 12)
        c.line(40, y - 3, W - 40, y - 3)
        y -= 18

    def two(left, lf, rtext, rf, size=10):
        nonlocal y
        text(c, 46, y, left, lf, size)
        right(c, W - 46, y, rtext, rf, size)
        y -= 13

    def bullets(items):
        nonlocal y
        for it in items:
            text(c, 54, y, "•", R, 9)
            text(c, 64, y, it, R, 9)
            y -= 12
        y -= 5

    heading("EDUCATION")
    two("Indian Institute of Technology Bombay", B, "Mumbai, India", R)
    two("Bachelor of Technology in Electrical Engineering", I, "Aug. 2021 – May 2025", I)
    y -= 6
    heading("EXPERIENCE")
    two("Software Engineering Intern", B, "May 2024 – July 2024", R)
    two("Swiggy", I, "Bengaluru, India", I)
    bullets(["Cut p95 latency of the restaurant search API from 420 ms to 180 ms by adding a Redis cache.",
             "Wrote load tests in Locust that now run on every release."])
    two("Undergraduate Researcher", B, "Jan. 2023 – Dec. 2023", R)
    two("IIT Bombay", I, "Mumbai, India", I)
    bullets(["Trained a CNN to classify power-grid faults with 94% accuracy on field data."])
    heading("PROJECTS")
    text(c, 46, y, "Gitty", B, 10)
    text(c, 46 + c.stringWidth("Gitty ", B, 10), y, "| Python, Flask, React, PostgreSQL", I, 10)
    right(c, W - 46, y, "June 2023 – Present", R, 10)
    y -= 13
    bullets(["A Git analytics dashboard used by 300+ students to track their open-source contributions."])
    heading("TECHNICAL SKILLS")
    for cat, items in [("Languages:", "Java, Python, C/C++, SQL, JavaScript"), ("Frameworks:", "React, Node.js, Flask, FastAPI"),
                       ("Developer Tools:", "Git, Docker, AWS")]:
        text(c, 46, y, cat, B, 9.5)
        text(c, 46 + c.stringWidth(cat + " ", B, 9.5), y, items, R, 9.5)
        y -= 12
    c.save()


# ------------------------------------------------------------------ 3. classic one column
def classic():
    c = canvas.Canvas(str(OUT / "classic.pdf"), pagesize=A4)
    y = 800
    for s, f, size in [("Priya Nair", B, 20), ("Kochi, India | priya.nair@example.com | +91 94470 11223", R, 9),
                       ("linkedin.com/in/priyanair | github.com/priya-nair", R, 9)]:
        text(c, 40, y, s, f, size)
        y -= 18
    sections = [
        ("SUMMARY", ["Computer science student who enjoys building clean, fast web apps."]),
        ("EXPERIENCE", ["Frontend Developer Intern | UST Global | Jan 2025 - Apr 2025",
                        "- Rebuilt the claims portal in React and TypeScript, cutting page load time by 35%.",
                        "- Wrote 120 unit tests with Jest.",
                        "Teaching Assistant | Model Engineering College | Aug 2024 - Dec 2024",
                        "- Ran weekly lab sessions on data structures for 60 students."]),
        ("EDUCATION", ["B.Tech in Computer Science, Model Engineering College, 2021 - 2025, CGPA 8.4"]),
        ("PROJECTS", ["Bus Buddy: Live bus tracking app for Kochi with Flutter and Firebase, used by 1,200 commuters.",
                      "Notes AI: Chrome extension that summarises lecture notes with an open-source LLM."]),
        ("SKILLS", ["JavaScript, TypeScript, React, Node.js, Flutter, Firebase, Git"]),
        ("CERTIFICATIONS", ["AWS Certified Cloud Practitioner (2024)"]),
        ("ACHIEVEMENTS", ["- Winner, Kerala Startup Mission Hackathon 2024"]),
    ]
    for head, lines in sections:
        y -= 6
        text(c, 40, y, head, B, 11)
        y -= 15
        for line in lines:
            text(c, 40, y, line, R, 9.5)
            y -= 13
    c.save()


# ------------------------------------------------------------------ 4. two-column sidebar
def two_column():
    c = canvas.Canvas(str(OUT / "two-column.pdf"), pagesize=A4)
    text(c, 30, 800, "Meera Joshi", B, 24)
    text(c, 30, 782, "Data Analyst", R, 11)
    left = [("CONTACT", B, 10), ("meera.joshi@example.com", R, 8.5), ("+91 99300 44556", R, 8.5), ("Pune, India", R, 8.5), ("", R, 8),
            ("SKILLS", B, 10), ("SQL", R, 8.5), ("Python", R, 8.5), ("Power BI", R, 8.5), ("Excel", R, 8.5), ("Statistics", R, 8.5), ("", R, 8),
            ("EDUCATION", B, 10), ("B.Sc Computer Science", B, 8.5), ("St. Xavier's College, Mumbai", R, 8.5), ("2020 – 2023", R, 8.5)]
    right_col = [("PROFILE", B, 10), ("Analyst who turns messy operational data into decisions. Three years", R, 9),
                 ("of SQL and dashboards for consulting clients.", R, 9), ("", R, 8),
                 ("WORK EXPERIENCE", B, 10), ("Data Analyst", B, 9.5), ("Deloitte India", R, 9), ("Jul 2023 – Present", R, 8.5),
                 ("• Built Power BI dashboards for 4 retail clients, used weekly by 40 managers.", R, 9),
                 ("• Reduced a monthly reconciliation from 3 days to 4 hours with Python.", R, 9), ("", R, 8),
                 ("PROJECTS", B, 10), ("Ride Demand Forecast", B, 9.5),
                 ("• Forecast hourly cab demand in Pune with a 9% error using open city data.", R, 9)]
    for col_x, rows in [(30, left), (210, right_col)]:
        y = 750
        for s, f, size in rows:
            if s:
                text(c, col_x, y, s, f, size)
            y -= 14
    c.save()


# ------------------------------------------------------------------ 5. Word document, experienced
def word():
    d = Document()
    d.add_paragraph("Vikram Rao")
    d.add_paragraph("Senior Data Analyst | vikram.rao@example.com | +91 98450 55555 | linkedin.com/in/vikramrao")
    d.add_paragraph("PROFESSIONAL SUMMARY")
    d.add_paragraph("Data analyst with eight years in e-commerce and analytics consulting.")
    d.add_paragraph("PROFESSIONAL EXPERIENCE")
    for title, company, dates, bullets in [
        ("Senior Data Analyst", "Flipkart, Bengaluru", "March 2021 – Present",
         ["Own the pricing experiments dashboard used by 25 category managers.", "Designed an A/B test framework in Python that cut analysis time by half."]),
        ("Data Analyst", "Mu Sigma, Bengaluru", "June 2018 – February 2021",
         ["Built demand forecasts in R for a US retail client.", "Automated 30 weekly SQL reports."]),
    ]:
        d.add_paragraph(title)
        d.add_paragraph(company)
        d.add_paragraph(dates)
        for b in bullets:
            d.add_paragraph(b, style="List Bullet")
    d.add_paragraph("EDUCATION")
    d.add_paragraph("MBA, Business Analytics, NMIMS Mumbai, 2016 – 2018")
    d.add_paragraph("B.Com, Christ University, 2013 – 2016")
    d.add_paragraph("KEY SKILLS")
    d.add_paragraph("SQL • Python • R • Tableau • Power BI • Excel • Statistics")
    d.save(str(OUT / "experienced.docx"))


# ------------------------------------------------------------------ 6. plain text, fresher
def plain():
    (OUT / "fresher.txt").write_text("""Kabir Singh
New Delhi | kabir.singh@example.com | 98110 22334 | github.com/kabirsingh

CAREER OBJECTIVE
Fresher looking for a web development role where I can learn from a strong team.

INTERNSHIPS
Web Development Intern — Infosys (Jan 2025 - Mar 2025)
* Built 6 responsive pages for an internal HR portal in React.
* Fixed 40 accessibility issues flagged by Lighthouse.

ACADEMIC DETAILS
B.Tech (CSE), Amity University, 2021 - 2025, CGPA 7.8
Class XII (CBSE), Delhi Public School, 2021, 91%
Class X (CBSE), Delhi Public School, 2019, 94%

PROJECTS
Expense Tracker: React Native app that splits shared expenses for hostel roommates.

SKILLS
HTML, CSS, JavaScript, React, Node.js, MongoDB
""", encoding="utf-8")


if __name__ == "__main__":
    modern(); jakes(); classic(); two_column(); word(); plain()
    print("written:", sorted(p.name for p in OUT.iterdir()))
